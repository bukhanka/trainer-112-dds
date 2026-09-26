#!/usr/bin/env python3
"""Convert the customer's incident classifier (xlsx) into data/classifier.json.

Usage:
    python3 scripts/convert-classifier.py <main.xlsx> [<old-with-scenarios.xlsx>] [-o data/classifier.json]

main.xlsx  — classifier v046_24: one sheet, 3 header rows, columns A–CU.
old.xlsx   — classifier v046_11: only its «Сценарий реагирования» column is used,
             matched to the main file by the final incident type (column K).

The xlsx files are not stored in the repository; the resulting JSON is.
Requires Python 3.9+ and openpyxl.
"""
from __future__ import annotations

import argparse
import collections
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

FIRST_ROUTE_COL = 14  # N
HIDDEN_MARK = "не отображается оператору 112"
NO_REACTION = "нет реагирования"

# Organization (row 1) → routing key. Matched as a lowercase substring, first hit wins.
ORG_KEYS: list[tuple[str, str]] = [
    ("классификатор мвд", "mvd"),
    ("классификатор смп", "smp"),
    ("классификатор мосгаз", "mosgaz"),
    ("цэмп", "cemp"),
    ("классификатор фсб", "fsb"),
    ("мособлгаз", "mosoblgaz"),
    ("автомобильные дороги ао", "okrug_roads"),
    ("автомобильные дороги", "autoroads"),
    ("мосгортранс", "mosgortrans"),
    ("гор. хозяйство", "city_economy"),
    ("гормост", "gormost"),
    ("канал имени москвы", "canal"),
    ("мгтс", "mgts"),
    ("метро", "metro"),
    ("мосводоканал", "mosvodokanal"),
    ("моэк", "moek"),
    ("моэск", "moesk"),
    ("оэк", "oek"),
    ("мослифт", "moslift"),
    ("цодд", "codd"),
    ("деп. жкх", "gkh_dep"),
    ("рбипк", "mosbez"),
    ("аппарат мэра", "mayor_office"),
    ("москоллектор", "moscollector"),
    ("ржд", "rzd"),
    ("департамент образования", "dep_education"),
    ("центррегионводхоз", "water_basin"),
    ("военная комендатура", "military_command"),
    ("оати", "oati"),
    ("мосводосток", "mosvodostok"),
    ("департамент ппиоос", "dep_nature"),
    ("департамент тсзн", "dep_social"),
    ("рсво", "rsvo"),
    ("эважд", "evazhd"),
    ("мсппн", "msppn"),
    ("дту_р", "dtu_ritual"),
    ("дту", "dtu"),
    ("росгвардия", "rosgvardia"),
    ("территориальные оив    тинао", "territorial_tinao"),
    ("территориальные оив тинао", "territorial_tinao"),
    ("территориальные оив", "territorial"),
    ("департамент строительства", "dep_construction"),
    ("комитет ветеринарии", "vet"),
    ("мосжилинспекция", "housing_inspection"),
    ("департамент культуры", "dep_culture"),
    ("гку цса", "csa"),
    ("гку нту", "ntu"),
    ("фсо", "fso"),
    ("гуп мср", "gup_msr"),
    ("комитет по туризму", "tourism"),
    ("дгп", "dgp"),
    ("цукб.бпла", "cukb_uav"),
    ("цукб", "cukb_mo"),
    ("организатор перевозок", "transport_organizer"),
    ("мосэкомониторинг", "ecomonitoring"),
    ("рхбз", "mo_rhbz"),
    ("ситиэнерго", "cityenergo"),
    ("гражданского строительства", "dep_civil_construction"),
]

# Inside «Классификатор МЧС» the second header row names separate organizations.
MCHS_SUBGROUPS = [("служба 101", "s101"), ("одс псц", "ods_psc"), ("мгпсс", "mgpss")]

# Second-row sub-channels that are not conditions: they get their own key suffix.
CHANNEL_SUFFIX = [
    ("дежурная служба", ""),  # «Дежурная служба АРМ-112» of РБиПК is the plate itself
    ("мкп, аналитика", "_analytics"),
    ("(куб)", "_kub"),
    ("пожары", "_fire"),
    ("интеграция", "_integration"),
    ("арм-112", "_arm"),
    ("по полигонам", "_polygons"),
    ("москва", "_moscow"),
]

# Condition text (rows 2–3) → condition code. Order matters: specific phrases first.
CONDITION_PATTERNS: list[tuple[str, str]] = [
    (r"не на месте", "victims_absent"),
    (r"(не выбран|признаки не выбраны|другие признаки)", "default"),
    (r"нет доступа", "no_access"),
    (r"угроза людям", "threat"),
    (r"постр", "victims"),
    (r"правонарушение", "offense"),
    (r"газификац", "gas"),
    (r"мед\. помощь", "med"),
    (r"эвакуац", "evac"),
    (r">5 чел", "crowd"),
    (r"перекрытие движени", "traffic"),
    (r"тоннель", "tunnel"),
    (r"^пеш", "crossing"),
    (r"^ав$", "road"),
    (r"реагирование всегда", "always"),
    (r"объектах связи", "comm"),
    (r"стройка", "construction"),
    (r"объект из перечня", "object_list"),
]

CARD_RE = re.compile(r"^карт\w*\s*-\s*1\d\d$", re.IGNORECASE)


def clean(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return text or None


def norm_key(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").lower().replace("ё", "е")).strip()


def normalize_label(value) -> str | None:
    """Routing cell → label; None when there is no route."""
    text = clean(value)
    if not text or text.lower() == NO_REACTION:
        return None
    if CARD_RE.match(text):  # «Карточка-112», «карточка -112», «Картчока-112», «Карточка-122» …
        return "карточка-112"
    return text


def merged_value(ws, row: int, col: int):
    """Value of a cell, looking through merged ranges (the value sits in the top-left cell)."""
    cell = ws.cell(row=row, column=col)
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return cell.value


def parse_route_columns(ws) -> list[dict]:
    columns = []
    org = None
    for col in range(FIRST_ROUTE_COL, ws.max_column + 1):
        row1 = clean(merged_value(ws, 1, col))
        if row1:
            org = row1
        # Row 2 is merged across columns only where the subgroup spans several of them.
        row2 = clean(merged_value(ws, 2, col))
        row3 = clean(ws.cell(row=3, column=col).value)
        org_l = norm_key(org)
        sub_l = norm_key(row2)

        key = None
        if "классификатор мчс" in org_l:
            key = next((k for pat, k in MCHS_SUBGROUPS if pat in sub_l), None)
        else:
            key = next((k for pat, k in ORG_KEYS if pat in org_l), None)
        if key is None:
            raise SystemExit(f"Unknown routing column {get_column_letter(col)}: {org!r} / {row2!r} / {row3!r}")

        cond_text = row3 or row2 or ""
        cond_l = norm_key(cond_text)
        condition = "always"
        if cond_text and not (row2 and not row3 and "классификатор мчс" in org_l):
            for pattern, code in CONDITION_PATTERNS:
                if re.search(pattern, cond_l):
                    condition = code
                    break
            else:
                # Not a condition: a sub-channel of the same organization (МКП, КУБ, интеграция…).
                suffix = next((s for pat, s in CHANNEL_SUFFIX if pat in cond_l), None)
                if suffix is None:
                    raise SystemExit(f"Unknown condition in {get_column_letter(col)}: {cond_text!r}")
                key += suffix
        columns.append(
            {
                "col": get_column_letter(col),
                "index": col - 1,
                "org": org,
                "subgroup": row2,
                "conditionText": row3,
                "routeKey": key,
                "condition": condition,
            }
        )

    seen = collections.Counter((c["routeKey"], c["condition"]) for c in columns)
    dup = [k for k, n in seen.items() if n > 1]
    if dup:
        raise SystemExit(f"Two columns map to the same (routeKey, condition): {dup}")
    return columns


QUESTION_START = re.compile(r"^(уточн\w+ .{8,}|где находится|знают ли|есть ли|что рядом)", re.IGNORECASE)


def split_questions(extra: str | None) -> list[str]:
    """Operator questions from column J: parts after «В:», or a direct instruction to clarify."""
    if not extra:
        return []
    if re.search(r"В\s*:", extra):
        parts = re.split(r"В\s*:", extra)[1:]
        out = []
        for part in parts:
            q = part.strip(" ;,.")
            if q and q != "(подтипы)":
                out.append(q)
        return out
    text = extra.strip()
    if QUESTION_START.match(text) or text.endswith("?"):
        return [text]
    return []


def load_scenarios(path: Path | None) -> tuple[dict[str, str], dict[int, str]]:
    """Final type → «Сценарий реагирования» (and code → scenario as a fallback) from the older file."""
    if path is None:
        return {}, {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [clean(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    try:
        scen_idx = next(i for i, h in enumerate(header) if h and "сценарий реагирования" in h.lower())
    except StopIteration:
        raise SystemExit(f"No «Сценарий реагирования» column in {path}")
    by_type: dict[str, str] = {}
    by_code: dict[int, str] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        final_type = clean(row[10]) if len(row) > 10 else None
        scen = clean(row[scen_idx]) if len(row) > scen_idx else None
        if not final_type or not scen:
            continue
        by_type.setdefault(norm_key(final_type), scen)
        if isinstance(row[4], (int, float)):
            by_code.setdefault(int(row[4]), scen)
    return by_type, by_code


def write_json_lines(path: Path, data: dict) -> None:
    """JSON with one list item per line: readable diffs when the classifier is updated."""
    def dumps(value) -> str:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

    parts = []
    for key, value in data.items():
        if isinstance(value, list):
            items = ",\n".join(dumps(v) for v in value)
            parts.append(f"{dumps(key)}:[\n{items}\n]")
        else:
            parts.append(f"{dumps(key)}:{dumps(value)}")
    path.write_text("{" + ",\n".join(parts) + "}\n", encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("main_xlsx", type=Path)
    ap.add_argument("scenarios_xlsx", type=Path, nargs="?")
    ap.add_argument("-o", "--out", type=Path, default=Path(__file__).resolve().parent.parent / "data" / "classifier.json")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.main_xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    columns = parse_route_columns(ws)
    scen_by_type, scen_by_code = load_scenarios(args.scenarios_xlsx)

    groups: list[dict] = []
    types: list[dict] = []
    group_by_id: dict[int, dict] = {}
    pending_group_name: str | None = None
    subgroup: str | None = None
    current_group: int | None = None
    scenario_hits = 0

    for row in ws.iter_rows(min_row=4, values_only=True):
        final_type = clean(row[10])
        if not final_type:
            name = clean(row[5])
            if name:  # group header row: «1 | Пожары и задымления», the last one has no number
                pending_group_name = name
            continue
        group_id = int(row[0])
        if group_id not in group_by_id:
            group = {"id": group_id, "name": pending_group_name or f"Группа {group_id}", "count": 0}
            group_by_id[group_id] = group
            groups.append(group)
        if group_id != current_group:
            current_group = group_id
            subgroup = None
        if clean(row[5]):
            subgroup = clean(row[5])
        group_by_id[group_id]["count"] += 1

        code = int(row[4])
        sign1, sign2, sign3 = clean(row[6]), clean(row[7]), clean(row[8])
        hidden = bool(sign1 and HIDDEN_MARK in sign1.lower())
        if hidden:
            sign1 = None
        extra = clean(row[9])
        scenario = scen_by_type.get(norm_key(final_type)) or scen_by_code.get(code)
        if scenario:
            scenario_hits += 1
        routes = []
        for c in columns:
            label = normalize_label(row[c["index"]] if c["index"] < len(row) else None)
            if label:
                routes.append([c["routeKey"], c["condition"], label])
        main_service = clean(row[12])
        types.append(
            {
                "code": code,
                "groupId": group_id,
                "subgroup": subgroup,
                "sign1": sign1,
                "sign2": sign2,
                "sign3": sign3,
                "extra": extra,
                "questions": split_questions(extra),
                "finalType": final_type,
                "ekpType": clean(row[11]),
                "mainService": main_service,
                "scenarioCode": scenario,
                "hiddenFromOperator": hidden,
                "routes": routes,
            }
        )

    out = {
        "source": {
            "file": args.main_xlsx.name,
            "version": "v046_24",
            "scenariosFrom": args.scenarios_xlsx.name if args.scenarios_xlsx else None,
            "note": "routes: [routeKey, condition, label]; label «карточка-112» = the service receives the 112 card as is",
        },
        "columns": [{k: v for k, v in c.items() if k != "index"} for c in columns],
        "groups": groups,
        "types": types,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    write_json_lines(args.out, out)

    n_routes = sum(len(t["routes"]) for t in types)
    print(
        f"{args.out}: {len(groups)} groups, {len(types)} types, {n_routes} routes, "
        f"{sum(t['hiddenFromOperator'] for t in types)} hidden, {scenario_hits} with scenario code, "
        f"{sum(bool(t['questions']) for t in types)} with questions"
    )


if __name__ == "__main__":
    main()
